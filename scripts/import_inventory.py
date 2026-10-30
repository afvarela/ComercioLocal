"""Importa el inventario inicial de Tienda El Descanso desde Excel."""
from __future__ import annotations

import argparse
import re
import sqlite3
import sys
from dataclasses import dataclass
from datetime import datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError as error:
    raise SystemExit("Falta openpyxl. Ejecute: py -3 -m pip install -r requirements.txt") from error

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_FILE = ROOT / "importaciones" / "INVENTARIO TIENDA EL DESCANSO.xlsx"
DEFAULT_DATABASE = ROOT / "data" / "comercio.sqlite3"
SKU_PREFIX = "TDESC-"
LOT_PREFIX = "TDESC-IMPORT-"
IMPORT_REFERENCE = "Importacion Excel Tienda El Descanso"
EXPECTED_HEADERS = ("nombre del producto", "um", "precio de venta", "costo", "inventario")
EXCLUDED_IMPORTED_NAMES = {
    "TOMATE CHONTO", "PAPA", "PLATANO MAD", "CEBOLLA CABEZONA", "AJO", "AGUACATE",
}


@dataclass(frozen=True)
class ProductRow:
    excel_row: int
    sku: str
    name: str
    source_unit: str
    unit: str
    price_basis: int
    price: int
    cost: int
    qty: int


def clean(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def number(value: object, field: str, row: int) -> Decimal:
    if value is None or clean(value) == "":
        return Decimal(0)
    try:
        result = Decimal(str(value).replace(",", "").strip())
    except (InvalidOperation, ValueError) as error:
        raise ValueError(f"Fila {row}: {field} no es numerico: {value!r}") from error
    if not result.is_finite() or result < 0:
        raise ValueError(f"Fila {row}: {field} debe ser mayor o igual a cero")
    return result


def money(value: object, field: str, row: int) -> int:
    return int((number(value, field, row) * 100).quantize(Decimal("1"), rounding=ROUND_HALF_UP))


def unit_and_qty(unit_value: object, qty_value: object, row: int) -> tuple[str, int, int]:
    source = clean(unit_value).lower()
    qty = number(qty_value, "inventario", row)
    if source in {"libra", "libras", "lb"}:
        return "g", 500, int((qty * 500).quantize(Decimal("1"), rounding=ROUND_HALF_UP))
    # "ml" describe el contenido del empaque; se vende como una unidad completa.
    return "und", 1000, int(qty.quantize(Decimal("1"), rounding=ROUND_HALF_UP))


def load_products(path: Path, sheet_name: str = "Productos") -> list[ProductRow]:
    if not path.exists():
        raise FileNotFoundError(f"No existe el archivo: {path}")
    workbook = load_workbook(path, read_only=True, data_only=True)
    if sheet_name not in workbook.sheetnames:
        raise ValueError(f"No existe la hoja {sheet_name!r}")
    sheet = workbook[sheet_name]
    rows = sheet.iter_rows(min_col=1, max_col=5, values_only=True)
    headers = tuple(clean(value).lower() for value in next(rows))
    if headers != EXPECTED_HEADERS:
        raise ValueError(
            "Las primeras cinco columnas deben ser: Nombre del Producto, UM, "
            f"Precio de venta, costo e inventario. Se encontraron: {headers}"
        )
    products = []
    for row, values in enumerate(rows, start=2):
        name = clean(values[0])
        if not name:
            continue
        if name.upper() in EXCLUDED_IMPORTED_NAMES:
            continue
        source_unit = clean(values[1]) or "un"
        unit, basis, qty = unit_and_qty(source_unit, values[4], row)
        products.append(ProductRow(
            row, f"{SKU_PREFIX}{row:05d}", name[:200], source_unit, unit, basis,
            money(values[2], "precio de venta", row), money(values[3], "costo", row), qty,
        ))
    workbook.close()
    if not products:
        raise ValueError("No se encontraron productos para importar")
    return products


def sql_value(value: object) -> str:
    return "NULL" if value is None else "'" + str(value).replace("'", "''") + "'"


def sql_statements(products: list[ProductRow]) -> str:
    lines = [
        "-- Importacion Tienda El Descanso, compatible con Cloudflare D1 (sin BEGIN/COMMIT).",
        f"DELETE FROM stock_moves WHERE product_id IN (SELECT id FROM products WHERE business_id=1 AND sku LIKE '{SKU_PREFIX}%') AND product_id NOT IN (SELECT product_id FROM sale_lines) AND product_id NOT IN (SELECT product_id FROM purchase_lines);",
        f"DELETE FROM lots WHERE business_id=1 AND product_id IN (SELECT id FROM products WHERE business_id=1 AND sku LIKE '{SKU_PREFIX}%') AND product_id NOT IN (SELECT product_id FROM sale_lines) AND product_id NOT IN (SELECT product_id FROM purchase_lines);",
        f"DELETE FROM products WHERE business_id=1 AND sku LIKE '{SKU_PREFIX}%' AND id NOT IN (SELECT product_id FROM sale_lines) AND id NOT IN (SELECT product_id FROM purchase_lines);",
        f"UPDATE products SET active=0 WHERE business_id=1 AND sku LIKE '{SKU_PREFIX}%';",
    ]
    created = datetime.now().astimezone().isoformat(timespec="seconds")
    for item in products:
        description = f"Importado desde Excel, fila {item.excel_row}. UM original: {item.source_unit}"
        lines.append(
            "INSERT INTO products(business_id,sku,barcode,name,description,category,unit,price,price_basis,image,min_stock,perishable,active,current_cost,max_stock,reorder_point,lead_time_days) VALUES("
            f"1,{sql_value(item.sku)},NULL,{sql_value(item.name)},{sql_value(description)},'Importacion Excel',{sql_value(item.unit)},{item.price},{item.price_basis},'',0,0,1,{item.cost},0,0,3) "
            "ON CONFLICT(business_id,sku) DO UPDATE SET name=excluded.name,description=excluded.description,category=excluded.category,unit=excluded.unit,price=excluded.price,price_basis=excluded.price_basis,active=1,current_cost=excluded.current_cost;"
        )
        if item.qty <= 0:
            continue
        code = f"{LOT_PREFIX}{item.excel_row:05d}"
        lines.append(
            "INSERT INTO lots(business_id,product_id,code,expires,qty,cost,cost_basis,supplier,reference,created,original_qty,purchase_id,status) "
            f"SELECT 1,id,{sql_value(code)},NULL,{item.qty},{item.cost},{item.price_basis},'Inventario inicial',{sql_value(IMPORT_REFERENCE)},{sql_value(created)},{item.qty},NULL,'Activo' FROM products "
            f"WHERE business_id=1 AND sku={sql_value(item.sku)} AND NOT EXISTS(SELECT 1 FROM sale_lines WHERE product_id=products.id) AND NOT EXISTS(SELECT 1 FROM lots WHERE code={sql_value(code)} AND product_id=products.id);"
        )
        lines.append(
            "INSERT INTO stock_moves(business_id,product_id,lot_id,user_id,created,qty,kind,reference,cost,cost_basis,balance_after,note) "
            f"SELECT 1,p.id,l.id,(SELECT id FROM users WHERE active=1 ORDER BY CASE WHEN role='admin' THEN 0 ELSE 1 END,id LIMIT 1),{sql_value(created)},{item.qty},'Inventario inicial',{sql_value(code)},{item.cost},{item.price_basis},{item.qty},{sql_value(IMPORT_REFERENCE)} "
            f"FROM products p JOIN lots l ON l.product_id=p.id AND l.code={sql_value(code)} WHERE p.business_id=1 AND p.sku={sql_value(item.sku)} AND NOT EXISTS(SELECT 1 FROM stock_moves WHERE lot_id=l.id AND reference={sql_value(code)});"
        )
    lines += [
        f"SELECT COUNT(*) productos_importados FROM products WHERE business_id=1 AND active=1 AND sku LIKE '{SKU_PREFIX}%';",
        f"SELECT COALESCE(SUM(qty),0) existencia_importada FROM lots WHERE business_id=1 AND product_id IN (SELECT id FROM products WHERE business_id=1 AND sku LIKE '{SKU_PREFIX}%');",
    ]
    return "\n".join(lines) + "\n"


def import_local(products: list[ProductRow], database: Path) -> tuple[Path, int, int]:
    sys.path.insert(0, str(ROOT))
    from server import Store
    Store(database)
    backup_dir = ROOT / "backups" / "inventario"
    backup_dir.mkdir(parents=True, exist_ok=True)
    backup = backup_dir / f"antes-importacion-{datetime.now():%Y%m%d-%H%M%S}.sqlite3"
    with sqlite3.connect(database) as source, sqlite3.connect(backup) as target:
        source.backup(target)
    created = datetime.now().astimezone().isoformat(timespec="seconds")
    with sqlite3.connect(database) as db:
        db.execute("PRAGMA foreign_keys=ON")
        user = db.execute(
            "SELECT id FROM users WHERE active=1 "
            "ORDER BY CASE WHEN role='admin' THEN 0 ELSE 1 END,id LIMIT 1"
        ).fetchone()
        if not user:
            raise RuntimeError("Primero abra la aplicacion y cree el usuario administrador.")
        used = db.execute(
            "SELECT "
            "(SELECT COUNT(*) FROM sale_lines WHERE product_id IN "
            "(SELECT id FROM products WHERE business_id=1 AND sku LIKE ?)) + "
            "(SELECT COUNT(*) FROM purchase_lines WHERE product_id IN "
            "(SELECT id FROM products WHERE business_id=1 AND sku LIKE ?))",
            (SKU_PREFIX + "%", SKU_PREFIX + "%"),
        ).fetchone()[0]
        if used:
            raise RuntimeError(
                "La importacion anterior ya tiene ventas o compras. "
                "Reinicie los datos de prueba antes de reemplazarla."
            )
        with db:
            ids = [
                row[0] for row in db.execute(
                    "SELECT id FROM products WHERE business_id=1 AND sku LIKE ?",
                    (SKU_PREFIX + "%",),
                )
            ]
            if ids:
                marks = ",".join("?" for _ in ids)
                db.execute(f"DELETE FROM stock_moves WHERE product_id IN ({marks})", ids)
                db.execute(f"DELETE FROM lots WHERE product_id IN ({marks})", ids)
                db.execute(f"DELETE FROM products WHERE id IN ({marks})", ids)
            for item in products:
                description = (
                    f"Importado desde Excel, fila {item.excel_row}. "
                    f"UM original: {item.source_unit}"
                )
                product_id = db.execute(
                    "INSERT INTO products("
                    "business_id,sku,barcode,name,description,category,unit,price,"
                    "price_basis,image,min_stock,perishable,active,current_cost,"
                    "max_stock,reorder_point,lead_time_days"
                    ") VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
                    (
                        1,item.sku,None,item.name,description,"Importacion Excel",
                        item.unit,item.price,item.price_basis,"",0,0,1,item.cost,0,0,3,
                    ),
                ).lastrowid
                if item.qty <= 0:
                    continue
                code = f"{LOT_PREFIX}{item.excel_row:05d}"
                lot_id = db.execute(
                    "INSERT INTO lots("
                    "business_id,product_id,code,expires,qty,cost,cost_basis,supplier,"
                    "reference,created,original_qty,purchase_id,status"
                    ") VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)",
                    (
                        1,product_id,code,None,item.qty,item.cost,item.price_basis,
                        "Inventario inicial",IMPORT_REFERENCE,created,item.qty,None,"Activo",
                    ),
                ).lastrowid
                db.execute(
                    "INSERT INTO stock_moves("
                    "business_id,product_id,lot_id,user_id,created,qty,kind,reference,"
                    "cost,cost_basis,balance_after,note"
                    ") VALUES(?,?,?,?,?,?,?,?,?,?,?,?)",
                    (
                        1,product_id,lot_id,user[0],created,item.qty,"Inventario inicial",
                        code,item.cost,item.price_basis,item.qty,IMPORT_REFERENCE,
                    ),
                )
    return backup, len(products), sum(item.qty > 0 for item in products)


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Importa el inventario inicial de Tienda El Descanso"
    )
    parser.add_argument("--file", type=Path, default=DEFAULT_FILE)
    parser.add_argument("--sheet", default="Productos")
    parser.add_argument(
        "--target", choices=("local", "sql", "validate"), default="validate"
    )
    parser.add_argument("--database", type=Path, default=DEFAULT_DATABASE)
    parser.add_argument(
        "--output", type=Path, default=ROOT / "tmp" / "import-inventory-d1.sql"
    )
    parser.add_argument("--confirm", default="")
    args = parser.parse_args()
    products = load_products(args.file.resolve(), args.sheet)
    print(
        f"Archivo validado: {len(products)} productos; "
        f"{sum(x.qty > 0 for x in products)} con existencia; "
        f"{sum(x.price == 0 for x in products)} con precio cero; "
        f"{sum(x.cost == 0 for x in products)} con costo cero."
    )
    if args.target == "validate":
        return
    if args.confirm != "IMPORTAR":
        raise SystemExit("Operacion cancelada. Use --confirm IMPORTAR.")
    if args.target == "sql":
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(
            sql_statements(products), encoding="utf-8", newline="\n"
        )
        print(f"SQL generado: {args.output.resolve()}")
        return
    backup, count, lots = import_local(products, args.database.resolve())
    print(
        f"Importacion local completada: {count} articulos "
        f"y {lots} lotes de inventario."
    )
    print(f"Respaldo anterior: {backup}")


if __name__ == "__main__":
    main()
